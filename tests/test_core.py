"""API 키 없이 돌아가는 단위·통합 테스트. 실제 판독 품질 테스트는 README의 '실전 테스트' 참고."""
import json
import os

import cv2
import pymupdf as fitz
import numpy as np
import pytest

from docread import tiling
from docread.consensus import apply_decisions, compare, count_markers, strip_markers
from docread.pdf_input import load_inputs
from docread.pipeline import Options, restructure, run
from docread.preprocess import estimate_skew, find_document_quad, prepare, rotate
from docread.profiles import EXAM, SAENGGIBU


# ---------- 전처리 ----------

def _ruled_page(w=1700, h=2200):
    img = np.full((h, w), 255, np.uint8)
    for y in range(200, h - 200, 80):
        cv2.line(img, (150, y), (w - 150, y), 0, 2)
        cv2.putText(img, "TEXT 1234 5678", (180, y - 20), cv2.FONT_HERSHEY_SIMPLEX, 1.2, 0, 2)
    return img


@pytest.mark.parametrize("angle", [-4.0, -1.5, 2.0, 6.0])
def test_deskew_recovers_angle(angle):
    tilted = rotate(_ruled_page(), angle)
    assert estimate_skew(tilted) == pytest.approx(angle, abs=0.3)
    assert abs(estimate_skew(rotate(tilted, -estimate_skew(tilted)))) < 0.3


@pytest.mark.parametrize("angle", [-3.0, 2.5])
def test_deskew_without_rule_lines(angle):
    img = np.full((2200, 1700), 255, np.uint8)
    for y in range(200, 2000, 70):
        cv2.putText(img, "plain paragraph text line 12345 abcde fghij", (150, y), cv2.FONT_HERSHEY_SIMPLEX, 1.3, 0, 3)
    assert estimate_skew(rotate(img, angle)) == pytest.approx(angle, abs=0.5)


def test_perspective_quad_found_on_photo():
    bg = np.full((1500, 1200, 3), 60, np.uint8)
    quad = np.array([[220, 160], [1010, 230], [1060, 1330], [150, 1280]], np.int32)
    cv2.fillConvexPoly(bg, quad, (245, 245, 245))
    found = find_document_quad(bg)
    assert found is not None
    for corner in quad:
        assert np.min(np.linalg.norm(found - corner, axis=1)) < 15


def test_prepare_outputs_work_width_and_flat_background():
    img = cv2.cvtColor(_ruled_page(), cv2.COLOR_GRAY2BGR)
    shade = np.linspace(0.45, 1.0, img.shape[1])[None, :, None]  # 왼쪽이 어두운 그림자
    img = (img * shade).astype(np.uint8)
    base, enh, meta = prepare(img, 2400, is_photo=False)
    assert base.shape[1] == enh.shape[1] == 2400
    # 그림자 쪽과 밝은 쪽 배경 밝기 차이가 크게 줄어야 한다
    left, right = enh[50:150, 50:300].mean(), enh[50:150, -300:-50].mean()
    assert abs(left - right) < 25
    assert meta["work_size"][0] == 2400


def test_strips_respect_model_limits():
    page = np.zeros((3394, 2400), np.uint8)
    strips = tiling.make_strips(page)
    assert len(strips) >= 2
    for s in strips:
        assert max(s.shape) <= 2576 and s.shape[0] * s.shape[1] <= 3_600_000
    covered = sum(s.shape[0] for s in strips) - (len(strips) - 1) * 90
    assert covered >= page.shape[0]


# ---------- 대조 ----------

def test_compare_ignores_spacing_and_table_syntax():
    a = "| 과목 | 등급 |\n|---|---|\n| 생활과 윤리 | 2 |"
    b = "과목 등급\n생활과윤리 2"
    c = compare(a, b)
    assert c.discrepancies == [] and not c.low_agreement


def test_compare_flags_digit_mismatch_and_apply():
    a = "| 생활과 윤리 | 3 | 85.3 |"
    b = "| 생활과 윤리 | 8 | 85.3 |"
    c = compare(a, b)
    assert len(c.discrepancies) == 1
    d = c.discrepancies[0]
    assert (d.a, d.b, d.digits) == ("3", "8", True)

    fixed, log = apply_decisions(a, c.discrepancies, {1: {"id": 1, "reading": "3", "confidence": "high", "reason": ""}})
    assert fixed == a and log[0]["status"] == "auto_resolved"

    unsure, log = apply_decisions(a, c.discrepancies, {1: {"id": 1, "reading": "8", "confidence": "medium", "reason": ""}})
    assert "⟦?8⟧" in unsure and count_markers(unsure) == 1 and strip_markers(unsure) == b

    missing, _ = apply_decisions(a, c.discrepancies, {})
    assert "⟦?3 ↔ 8⟧" in missing and strip_markers(missing) == a


def test_compare_handles_insertions():
    c = compare("가 나 라", "가 나 다 라")
    assert len(c.discrepancies) == 1 and c.discrepancies[0].a == "" and c.discrepancies[0].b == "다"
    fixed, _ = apply_decisions("가 나 라", c.discrepancies, {1: {"id": 1, "reading": " 다", "confidence": "high", "reason": ""}})
    assert fixed == "가 나 다 라"


def test_low_agreement_detected():
    assert compare("완전히 다른 내용 하나 둘 셋", "전혀 무관한 텍스트 넷 다섯 여섯").low_agreement


# ---------- 검증 규칙 ----------

def _row(**kw):
    base = dict(year="2", semester="1", category="일반", group="사회", subject="생활과 윤리", units="3",
                raw_score="92", average="70.1", std_dev="15.2", achievement="A", enrolled="152",
                rank_grade="2", distribution="", uncertain_fields=[])
    base.update(kw)
    return base


def test_saenggibu_rules():
    good = {"grade_rows": [_row()], "subject_notes": [], "activities": [], "behavior": [], "other_sections": []}
    assert [i for i in SAENGGIBU.validate(good) if i["level"] != "info"] == []

    bad = {"grade_rows": [
        _row(rank_grade="12"),                       # 등급 범위 밖
        _row(subject="한국지리", raw_score="78", achievement="A"),  # 원점수-성취도 불일치
        _row(subject="경제", category="진로선택", achievement="A", raw_score="", average="", std_dev="",
             rank_grade="", distribution="A(35.2) B(41.0) C(28.8)"),  # 분포 합계 105
        _row(subject="세계사", units="0", uncertain_fields=["rank_grade"]),
    ], "subject_notes": [{"year": "2", "subject": "경제", "text": "가" * 700, "uncertain_fields": []}],
        "activities": [], "behavior": [], "other_sections": []}
    msgs = " / ".join(i["message"] for i in SAENGGIBU.validate(bad))
    for frag in ("석차등급 '12'", "성취도 A 구간 불일치", "분포비율 합계", "단위수 '0'", "판독 불확실", "기재 한도"):
        assert frag in msgs, frag


def test_exam_rules():
    ch = lambda labels: [{"label": l, "text": f"선지{i}"} for i, l in enumerate(labels)]
    data = {"header": "", "questions": [
        {"number": "1", "stem": "옳은 것은?", "passage": "", "choices": ch("①②③④⑤"), "points": "", "figure_note": "", "uncertain_fields": []},
        {"number": "3", "stem": "고른 것은?", "passage": "<보기> ㄱ. 가 ㄴ. 나", "choices": [
            {"label": "①", "text": "ㄱ, ㄴ"}, {"label": "②", "text": "ㄱ, ㄷ"}, {"label": "③", "text": "ㄴ, ㄷ"},
            {"label": "④", "text": "ㄴ, ㄹ"}], "points": "", "figure_note": "", "uncertain_fields": ["stem"]},
    ]}
    msgs = " / ".join(i["message"] for i in EXAM.validate(data))
    for frag in ("빠진 문항 번호: [2]", "선지 4개", "ㄷ, ㄹ 진술", "판독 불확실"):
        assert frag in msgs, frag


# ---------- 파이프라인 (가짜 엔진) ----------

class FakeClaude:
    name = "claude"
    model = "claude-opus-5-5"

    def __init__(self):
        from docread.engines.base import Usage
        self.usage = Usage()
        self.adjudicated = []

    def transcribe(self, img, profile):
        self.usage.add(1000, 100)
        # 보정본(흑백 2차원)과 원본(컬러)에서 숫자 하나를 다르게 읽은 상황을 흉내 낸다.
        grade = "3" if img.ndim == 2 else "8"
        return f"| 학년 | 학기 | 과목 | 등급 |\n|---|---|---|---|\n| 2 | 1 | 생활과 윤리 | {grade} |"

    def adjudicate(self, img, items):
        self.adjudicated.append(items)
        return {d["id"]: {"id": d["id"], "reading": "3", "confidence": "high", "reason": "선명"} for d in items}

    def structure(self, transcript, profile):
        self.usage.add(500, 200)
        return {"grade_rows": [_row(rank_grade="3")], "subject_notes": [], "activities": [], "behavior": [], "other_sections": []}


def _make_pdfs(tmp_path):
    text_pdf = tmp_path / "text.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Student record text layer " * 5)
    doc.save(text_pdf)

    scan_pdf = tmp_path / "scan.pdf"
    png = tmp_path / "scan.png"
    cv2.imwrite(str(png), _ruled_page())
    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(page.rect, filename=str(png))
    page.insert_text((72, 72), "invisible scanner ocr layer " * 5, render_mode=3)
    doc.save(scan_pdf)
    return str(text_pdf), str(scan_pdf)


def test_text_layer_trust(tmp_path):
    text_pdf, scan_pdf = _make_pdfs(tmp_path)
    pages = load_inputs([text_pdf, scan_pdf], dpi=100)
    assert [p.kind for p in pages] == ["pdf_text", "pdf_scan"]
    assert "Student record" in pages[0].text_layer and pages[1].text_layer is None


def test_pipeline_end_to_end(tmp_path):
    text_pdf, scan_pdf = _make_pdfs(tmp_path)
    fake = FakeClaude()
    out = tmp_path / "out"
    res = run([text_pdf, scan_pdf], str(out), Options(dpi=100), claude=fake, second=fake, progress=lambda m: None)

    p1, p2 = res["pages"]
    assert p1["method"] == "text_layer"
    assert p2["method"] == "ocr_double_adjudicated"
    assert p2["discrepancies"][0]["a"] == "3" and p2["discrepancies"][0]["b"] == "8"
    assert "| 3 |" in p2["transcript"] and "⟦?" not in p2["transcript"]
    assert res["structured"]["grade_rows"][0]["rank_grade"] == "3"
    assert res["usage"]["calls"] == 3  # 전사 2 + 구조화 1 (판정은 가짜라 미집계)
    assert res["exports"] == ["text_전사.docx", "text_정리.docx"]
    for f in ("result.json", "transcript.md", "transcript_clean.md", "saenggibu.json", *res["exports"],
              "pages/p002_original.png", "pages/p002_enhanced.png"):
        assert (out / f).exists(), f

    # 사람이 고친 뒤 구조화만 다시
    saved = json.loads((out / "result.json").read_text(encoding="utf-8"))
    saved["pages"][1]["transcript"] += "\n⟦?7⟧"
    (out / "result.json").write_text(json.dumps(saved, ensure_ascii=False), encoding="utf-8")
    res2 = restructure(str(out), claude=fake, progress=lambda m: None)
    assert any("불확실 표시" in i["message"] for i in res2["issues"])


def test_page_failure_is_isolated(tmp_path):
    _, scan_pdf = _make_pdfs(tmp_path)

    class Boom(FakeClaude):
        def transcribe(self, img, profile):
            raise RuntimeError("network down")

    res = run([scan_pdf], str(tmp_path / "o"), Options(dpi=100, structure=False), claude=Boom(), second=None,
              progress=lambda m: None)
    assert res["pages"][0]["error"].startswith("RuntimeError")
    assert res["issues"][0]["level"] == "error"


def test_server_smoke(tmp_path, monkeypatch):
    monkeypatch.setenv("DOCREAD_DATA_DIR", str(tmp_path))
    import importlib

    import docread.config
    importlib.reload(docread.config)
    import webapp.server as server
    importlib.reload(server)
    from fastapi.testclient import TestClient

    client = TestClient(server.app)
    assert client.get("/").status_code == 200
    assert {p["key"] for p in client.get("/api/profiles").json()} >= {"saenggibu", "exam"}
    assert client.get("/api/jobs").json() == []
    assert client.get("/api/jobs/../../etc").status_code == 404


# ---------- 정확도 측정 ----------

def test_evaluate_scores():
    from docread.evaluate import score
    truth = "| 생활과 윤리 | 3 | 85.3 |"
    assert score(truth, truth)["cer"] == 0
    r = score("| 생활과 윤리 | ⟦?8⟧ | 85.3 |", truth)
    assert r["digit_error_rate"] > 0 and r["errors_flagged_ratio"] == 1.0 and r["flags"] == 1
    r = score("| 생활과 윤리 | 8 | 85.3 |", truth)
    assert r["errors_flagged_ratio"] == 0.0


# ---------- Word 변환 ----------

def _word_result(profile, structured):
    return {"meta": {"title": "테스트 문서", "profile": profile, "profile_label": "x", "page_count": 2, "inputs": ["a.pdf"]},
            "pages": [{"page": 1, "source": "a.pdf p.1",
                       "transcript": "# 교과학습발달상황\n| 과목 | 등급 |\n|---|---|\n| 생활과 윤리 | ⟦?3⟧ |\n**굵게** 본문 __밑줄__"},
                      {"page": 2, "source": "a.pdf p.2", "transcript": "둘째 쪽"}],
            "structured": structured, "issues": [{"level": "error", "where": "x", "message": "테스트 오류"}]}


def test_word_transcript_tables_and_highlights(tmp_path):
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docread.word import export_all

    files = export_all(_word_result("generic", None), str(tmp_path))
    assert files == ["테스트_문서_전사.docx"]
    d = Document(tmp_path / files[0])
    assert len(d.tables) == 1
    t = d.tables[0]
    assert [c.text for c in t.rows[0].cells] == ["과목", "등급"]
    run = t.cell(1, 1).paragraphs[0].runs[0]
    assert run.text == "3" and run.font.highlight_color == WD_COLOR_INDEX.YELLOW
    body = "\n".join(p.text for p in d.paragraphs)
    assert "굵게 본문 밑줄" in body and "둘째 쪽" in body and "⟦" not in body


def test_word_summary_docs(tmp_path):
    from docx import Document
    from docread.word import export_all

    sg = {"school_year_info": "", "grade_rows": [_row(uncertain_fields=["rank_grade"])],
          "subject_notes": [{"year": "2", "subject": "생활과 윤리", "text": "토론 활동에서…", "uncertain_fields": []}],
          "activities": [], "behavior": [], "other_sections": []}
    files = export_all(_word_result("saenggibu", sg), str(tmp_path))
    d = Document(tmp_path / files[1])
    assert d.tables[0].cell(1, 4).text == "생활과 윤리" and d.tables[0].cell(1, 11).text == "2"
    assert "w:fill=\"FFF2B3\"" in d.tables[0].cell(1, 11)._tc.xml
    assert any("토론 활동" in p.text for p in d.paragraphs)
    assert "테스트 오류" in d.tables[-1].cell(1, 2).text

    ex = {"header": "2학기 기말", "questions": [{"number": "1", "stem": "옳은 것은?", "passage": "<보기>\nㄱ. 가\n| a | b |\n|---|---|\n| 1 | 2 |",
                                                "choices": [{"label": "①", "text": "ㄱ"}], "points": "[3점]", "figure_note": "", "uncertain_fields": []}]}
    files = export_all(_word_result("exam", ex), str(tmp_path))
    d = Document(tmp_path / files[1])
    text = "\n".join(p.text for p in d.paragraphs)
    assert "1. 옳은 것은?" in text and "① ㄱ" in text
    box = d.tables[0].cell(0, 0)
    assert "ㄱ. 가" in box.text
    assert [c.text for c in box.tables[0].rows[1].cells] == ["1", "2"]
