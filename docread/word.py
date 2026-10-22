"""판독 결과를 Word(.docx)로 만든다.

- <제목>_전사.docx : 원문 그대로의 전사본. 마크다운 표는 실제 Word 표로, ⟦?⟧ 표시는 노란 형광펜으로.
- <제목>_정리.docx : 구조화 결과를 보기 좋게 정리한 문서(생기부: 성적표·세특·창체·행특 / 시험지: 문항 카드).
                     검증 결과(확인 필요 목록)를 부록으로 붙인다.
"""
from __future__ import annotations

import os
import re
import time

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "맑은 고딕"
GRAY = RGBColor(0x6A, 0x73, 0x7D)
UNSURE_FILL = "FFF2B3"
INLINE_RE = re.compile(r"(⟦\?[^⟧]*⟧|\*\*[^*]+\*\*|__[^_]+__)")
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def export_all(result: dict, out_dir: str, mark_uncertain: bool = True) -> list[str]:
    """결과 폴더에 Word 파일들을 만들고 상대 경로 목록을 돌려준다."""
    base = _safe_name(result["meta"].get("title") or os.path.splitext(result["meta"]["inputs"][0])[0])
    files = []
    path = f"{base}_전사.docx"
    transcript_doc(result, mark_uncertain).save(os.path.join(out_dir, path))
    files.append(path)
    if result.get("structured"):
        path = f"{base}_정리.docx"
        summary_doc(result).save(os.path.join(out_dir, path))
        files.append(path)
    return files


# ---------- 원문 전사본 ----------

def transcript_doc(result: dict, mark_uncertain: bool = True) -> Document:
    doc = _new_doc()
    _title(doc, result, "전사본")
    if mark_uncertain:
        p = doc.add_paragraph()
        r = p.add_run("노란 형광펜")
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        _small(p.add_run(" = 판독이 불확실해 원본 대조가 필요한 곳"))
    pages = result["pages"]
    for i, page in enumerate(pages):
        _small(doc.add_paragraph().add_run(f"— {page['page']}쪽 · {page.get('source', '')} —"))
        if page.get("error"):
            doc.add_paragraph(f"[판독 실패: {page['error']}]")
        add_markdown(doc, page.get("transcript", ""), mark_uncertain)
        if i < len(pages) - 1:
            doc.add_page_break()
    return doc


def add_markdown(doc: Document, text: str, mark_uncertain: bool = True) -> None:
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.lstrip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            _md_table(doc, block, mark_uncertain)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            h = doc.add_heading(level=min(len(m.group(1)), 3))
            _inline(h, m.group(2), mark_uncertain)
        elif line.strip():
            _inline(doc.add_paragraph(), line, mark_uncertain)
        i += 1


def _md_table(doc: Document, block: list[str], mark_uncertain: bool) -> None:
    rows = [_split_row(r) for r in block if not TABLE_SEP_RE.match(r)]
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = len(block) > 1 and TABLE_SEP_RE.match(block[1]) is not None
    for ri, row in enumerate(rows):
        for ci in range(ncol):
            val = row[ci] if ci < len(row) else ""
            cell = t.cell(ri, ci)
            p = cell.paragraphs[0]
            _inline(p, val, mark_uncertain, size=9)
            if header and ri == 0:
                for r in p.runs:
                    r.bold = True
                _shade(cell, "EEF1F5")
    doc.add_paragraph()


def _split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def _inline(par, text: str, mark_uncertain: bool, size: float | None = None) -> None:
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("⟦?"):
            inner = part[2:-1].split(" ↔ ")[0] or "?"
            r = par.add_run(inner)
            if mark_uncertain:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif part.startswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("__"):
            r = par.add_run(part[2:-2])
            r.underline = True
        else:
            r = par.add_run(part)
        if size:
            r.font.size = Pt(size)


# ---------- 정리본 ----------

def summary_doc(result: dict) -> Document:
    prof = result["meta"]["profile"]
    s = result["structured"]
    doc = _new_doc(landscape=(prof == "saenggibu"))
    _title(doc, result, "정리본")
    if prof == "saenggibu":
        _saenggibu(doc, s)
    elif prof == "exam":
        _exam(doc, s)
    _issues(doc, result.get("issues") or [])
    return doc


GRADE_COLS = [("year", "학년"), ("semester", "학기"), ("category", "구분"), ("group", "교과"), ("subject", "과목"),
              ("units", "단위"), ("raw_score", "원점수"), ("average", "평균"), ("std_dev", "표준편차"),
              ("achievement", "성취도"), ("enrolled", "수강자"), ("rank_grade", "등급"), ("distribution", "분포비율")]
GRADE_WIDTHS_CM = [1.1, 1.1, 1.8, 3.6, 4.0, 1.1, 1.5, 1.5, 1.7, 1.4, 1.5, 1.2, 4.4]


def _saenggibu(doc: Document, s: dict) -> None:
    if s.get("school_year_info"):
        _small(doc.add_paragraph().add_run(s["school_year_info"]))
    rows = s.get("grade_rows") or []
    doc.add_heading("1. 교과학습발달상황", level=1)
    if rows:
        t = doc.add_table(rows=len(rows) + 1, cols=len(GRADE_COLS))
        t.style = "Table Grid"
        t.autofit = False
        for ci, w in enumerate(GRADE_WIDTHS_CM):
            t.columns[ci].width = Cm(w)  # 표 격자(tblGrid) — LibreOffice·한글은 이 값을 본다
            for cell in t.columns[ci].cells:
                cell.width = Cm(w)  # 칸 너비 — Word는 이 값을 본다
        for ci, (_, label) in enumerate(GRADE_COLS):
            c = t.cell(0, ci)
            r = c.paragraphs[0].add_run(label)
            r.bold, r.font.size = True, Pt(8.5)
            _shade(c, "EEF1F5")
        for ri, row in enumerate(rows, 1):
            unsure = set(row.get("uncertain_fields") or [])
            for ci, (key, _) in enumerate(GRADE_COLS):
                c = t.cell(ri, ci)
                c.paragraphs[0].add_run(str(row.get(key, ""))).font.size = Pt(8.5)
                if key in unsure:
                    _shade(c, UNSURE_FILL)
        _small(doc.add_paragraph().add_run("노란 칸 = 판독 불확실, 원본 대조 필요"))
    else:
        doc.add_paragraph("(추출된 성적 없음)")

    for title, key, label_keys in (("2. 세부능력 및 특기사항", "subject_notes", ("year", "subject")),
                                   ("3. 창의적 체험활동", "activities", ("year", "area", "hours")),
                                   ("4. 행동특성 및 종합의견", "behavior", ("year",)),
                                   ("5. 기타 항목", "other_sections", ("title",))):
        items = s.get(key) or []
        if not items:
            continue
        doc.add_heading(title, level=1)
        for it in items:
            label = " · ".join(_label(k, it.get(k)) for k in label_keys if it.get(k))
            p = doc.add_paragraph()
            r = p.add_run(label)
            r.bold = True
            if it.get("uncertain_fields"):
                r2 = p.add_run("  [확인 필요]")
                r2.font.highlight_color = WD_COLOR_INDEX.YELLOW
            add_markdown(doc, it.get("text", ""))


def _label(key: str, v: str) -> str:
    v = str(v)
    if key == "year" and v.isdigit():
        return f"{v}학년"
    if key == "hours" and v:
        return f"{v}시간"
    return v


def _exam(doc: Document, s: dict) -> None:
    if s.get("header"):
        _small(doc.add_paragraph().add_run(s["header"]))
    for q in s.get("questions") or []:
        unsure = set(q.get("uncertain_fields") or [])
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        r = p.add_run(f"{q.get('number', '')}. ")
        r.bold = True
        r = p.add_run(q.get("stem", ""))
        if "stem" in unsure:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if q.get("points"):
            _small(p.add_run(f"  {q['points']}"))
        if q.get("passage"):
            box = doc.add_table(rows=1, cols=1)
            box.style = "Table Grid"
            cell = box.cell(0, 0)
            if "passage" in unsure:
                _shade(cell, UNSURE_FILL)
            cell.paragraphs[0].text = ""
            _cell_markdown(cell, q["passage"])
        if q.get("figure_note"):
            r = doc.add_paragraph().add_run(f"[그림] {q['figure_note']}")
            r.italic, r.font.color.rgb = True, GRAY
        for c in q.get("choices") or []:
            cp = doc.add_paragraph(f"{c.get('label', '')} {c.get('text', '')}")
            cp.paragraph_format.left_indent = Cm(0.6)
            cp.paragraph_format.space_after = Pt(0)
            if "choices" in unsure:
                for r in cp.runs:
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _cell_markdown(cell, text: str) -> None:
    """자료 박스(표 한 칸) 안에 문단과 표를 넣는다. 마크다운 표는 칸 안의 실제 Word 표가 된다."""
    lines = text.splitlines()
    first, i = True, 0
    while i < len(lines):
        if lines[i].lstrip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = [_split_row(r) for r in block if not TABLE_SEP_RE.match(r)]
            if rows:
                ncol = max(len(r) for r in rows)
                inner = cell.add_table(rows=len(rows), cols=ncol)
                inner.style = "Table Grid"
                inner.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(rows):
                    for ci in range(ncol):
                        _inline(inner.cell(ri, ci).paragraphs[0], row[ci] if ci < len(row) else "", True, size=9)
                first = False
            continue
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        _inline(p, lines[i], True, size=9.5)
        i += 1
    lead = cell.paragraphs[0]._p
    if not cell.paragraphs[0].text and lead.getnext() is not None:  # 표로 시작하면 빈 첫 문단 제거
        lead.getparent().remove(lead)
    if cell._tc[-1].tag == qn("w:tbl"):
        cell.add_paragraph()  # Word 규칙: 칸은 문단으로 끝나야 한다


def _issues(doc: Document, issues: list[dict]) -> None:
    doc.add_page_break()
    doc.add_heading("부록. 자동 검증 결과", level=1)
    if not issues:
        doc.add_paragraph("검증 규칙상 문제 없음")
        return
    t = doc.add_table(rows=len(issues) + 1, cols=3)
    t.style = "Table Grid"
    for ci, h in enumerate(("수준", "위치", "내용")):
        t.cell(0, ci).paragraphs[0].add_run(h).bold = True
    lv = {"error": "오류", "warning": "확인", "info": "참고"}
    for ri, it in enumerate(issues, 1):
        for ci, v in enumerate((lv.get(it["level"], it["level"]), it["where"], it["message"])):
            t.cell(ri, ci).paragraphs[0].add_run(v).font.size = Pt(9)
        if it["level"] == "error":
            _shade(t.cell(ri, 0), "F8D7DA")


# ---------- 공통 ----------

def _new_doc(landscape: bool = False) -> Document:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        hs = doc.styles[name]
        hs.font.name = FONT
        hs.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.8))
    return doc


def _title(doc: Document, result: dict, kind: str) -> None:
    meta = result["meta"]
    doc.add_heading(f"{meta.get('title') or meta.get('profile_label', '')} — {kind}", level=0)
    _small(doc.add_paragraph().add_run(
        f"{meta.get('profile_label', '')} · {meta.get('page_count')}쪽 · 원본: {', '.join(meta.get('inputs', []))} · "
        f"생성 {time.strftime('%Y-%m-%d %H:%M')}"))


def _small(run) -> None:
    run.font.size = Pt(8.5)
    run.font.color.rgb = GRAY


def _shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _safe_name(name: str) -> str:
    return re.sub(r"[^\w가-힣\-]+", "_", name).strip("_")[:60] or "판독결과"
